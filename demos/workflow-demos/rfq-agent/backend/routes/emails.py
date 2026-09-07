"""Email inbox routes — reads Gmail via IMAP and processes emails as RFQs."""
import asyncio
import logging
import os
import time
from fastapi import APIRouter, HTTPException, BackgroundTasks
from db.database import get_db
from sqlalchemy.ext.asyncio import AsyncSession
from fastapi import Depends

logger = logging.getLogger(__name__)

router = APIRouter(tags=["emails"])

# Simple in-memory cache: key → (timestamp, data)
_inbox_cache: dict[str, tuple[float, dict]] = {}
_CACHE_TTL = 30  # seconds


def _cache_key(account: str | None, folder: str, limit: int) -> str:
    return f"{account}:{folder}:{limit}"


@router.get("/emails")
async def list_emails(folder: str = "INBOX", limit: int = 25, account: str | None = None, db: AsyncSession = Depends(get_db)):
    logger.info(f"[INBOX] Fetching {limit} emails from {folder}")

    cache_key = _cache_key(account, folder, limit)
    cached = _inbox_cache.get(cache_key)
    if cached and (time.time() - cached[0]) < _CACHE_TTL:
        logger.info(f"[INBOX] Serving from cache (age {int(time.time() - cached[0])}s)")
        return cached[1]

    try:
        from services.gmail_inbox import list_emails as _list, list_emails_oauth
        from services.gmail_oauth import credentials_from_token_json
        from db.models import ConnectedAccount
        from sqlalchemy import select

        # Find the account to use
        query = select(ConnectedAccount).where(ConnectedAccount.active == True)
        if account:
            query = query.where(ConnectedAccount.email == account)
        else:
            query = query.limit(1)

        result = await db.execute(query)
        connected = result.scalars().first()

        if connected:
            creds = credentials_from_token_json(connected.token_json)
            if creds:
                try:
                    # REST API — fast, no IMAP connection overhead
                    from services.gmail_inbox import list_emails_rest
                    emails = await asyncio.to_thread(list_emails_rest, creds, folder, limit)
                    logger.info(f"[INBOX] Fetched {len(emails)} emails via REST for {connected.email}")
                except Exception as rest_err:
                    logger.warning(f"[INBOX] REST failed, falling back to IMAP: {rest_err}")
                    emails = await asyncio.to_thread(list_emails_oauth, connected.email, creds.token, folder, limit)
                    logger.info(f"[INBOX] Fetched {len(emails)} emails via IMAP for {connected.email}")
                response = {"emails": emails, "total": len(emails), "account": connected.email}
                _inbox_cache[cache_key] = (time.time(), response)
                return response

        # Fallback to app password
        emails = await asyncio.to_thread(_list, folder=folder, limit=limit)
        logger.info(f"[INBOX] Fetched {len(emails)} emails via app password")
        response = {"emails": emails, "total": len(emails)}
        _inbox_cache[cache_key] = (time.time(), response)
        return response
    except RuntimeError as e:
        raise HTTPException(status_code=503, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gmail connection failed: {e}")


@router.get("/emails/watcher/status")
async def watcher_status():
    """Return current email watcher status and connected accounts."""
    from services.email_watcher import _running
    from db.database import SessionLocal
    from db.models import ConnectedAccount
    from sqlalchemy import select
    interval = int(os.getenv("EMAIL_POLL_INTERVAL", 60))
    accounts = []
    try:
        async with SessionLocal() as db:
            result = await db.execute(select(ConnectedAccount).where(ConnectedAccount.active == True))
            accounts = [a.email for a in result.scalars().all()]
    except Exception:
        pass
    return {
        "running": _running,
        "poll_interval_seconds": interval,
        "watching": accounts,
    }


@router.post("/emails/watcher/trigger")
async def watcher_trigger():
    """Manually trigger an immediate inbox check across all connected accounts."""
    from services.email_watcher import _check_connected_accounts
    try:
        loop = asyncio.get_running_loop()
        await _check_connected_accounts(loop)
        return {"status": "ok", "message": "Inbox check triggered"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/emails/{email_id}")
async def get_email(email_id: str, account: str | None = None, db: AsyncSession = Depends(get_db)):
    try:
        from services.gmail_inbox import get_email as _get, get_email_oauth, get_thread_emails
        from services.gmail_oauth import credentials_from_token_json
        from db.models import ConnectedAccount
        from sqlalchemy import select

        # Try OAuth first
        query = select(ConnectedAccount).where(ConnectedAccount.active == True)
        if account:
            query = query.where(ConnectedAccount.email == account)
        else:
            query = query.limit(1)
        result = await db.execute(query)
        connected = result.scalars().first()
        creds = None

        if connected:
            creds = credentials_from_token_json(connected.token_json)
            if creds:
                if email_id.isdigit():
                    # Old IMAP seq number stored in DB — use IMAP path
                    data = await asyncio.to_thread(get_email_oauth, connected.email, creds.token, email_id)
                else:
                    # Hex Gmail REST message ID — use REST
                    from services.gmail_inbox import get_email_rest
                    data = await asyncio.to_thread(get_email_rest, creds, email_id)
            else:
                data = await asyncio.to_thread(_get, email_id)
        else:
            data = await asyncio.to_thread(_get, email_id)

        thread_id = data.get("thread_id")

        from db.models import RFQSubmission
        from sqlalchemy import select, or_

        thread_messages = []
        found_rfq = None

        def _thread_variants(tid: str) -> list[str]:
            """Return both hex and decimal variants of a Gmail thread/message ID."""
            variants = [tid]
            try:
                if tid.isdigit():
                    variants.append(hex(int(tid))[2:])
                else:
                    variants.append(str(int(tid, 16)))
            except Exception:
                pass
            return variants

        async def find_rfq() -> dict | None:
            """Find an RFQ linked to this email via thread_id or email_id (any ID format)."""
            nonlocal found_rfq
            queries = []
            if thread_id:
                queries.append(RFQSubmission.thread_id.in_(_thread_variants(thread_id)))
            queries.append(RFQSubmission.email_id.in_(_thread_variants(email_id)))
            result = await db.execute(
                select(RFQSubmission).where(or_(*queries)).limit(1)
            )
            rec = result.scalars().first()
            if rec:
                found_rfq = {
                    "rfq_id": rec.rfq_id,
                    "status": rec.status,
                    "total": rec.total,
                    "draft_quote": rec.draft_quote,
                }

        if thread_id:
            async def fetch_thread():
                nonlocal thread_messages
                if connected and creds and not email_id.isdigit():
                    from services.gmail_inbox import get_thread_emails_rest
                    thread_messages = await asyncio.to_thread(get_thread_emails_rest, creds, thread_id)
                else:
                    thread_messages = await asyncio.to_thread(get_thread_emails, thread_id)

            await asyncio.gather(find_rfq(), fetch_thread())
        else:
            thread_messages = [data]
            await find_rfq()

        any_thread_rfq = found_rfq

        for msg in thread_messages:
            msg["rfq"] = any_thread_rfq

        data["thread_messages"] = thread_messages
        data["rfq"] = any_thread_rfq

        return data
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/emails/{email_id}/process")
async def process_email_as_rfq(
    email_id: str,
    background_tasks: BackgroundTasks,
    db: AsyncSession = Depends(get_db),
):
    """Parse email body and run it through the RFQ pipeline."""
    # Guard: check if an RFQ has already been generated for this email message
    from db.models import RFQSubmission
    from sqlalchemy import select
    result = await db.execute(
        select(RFQSubmission).where(RFQSubmission.email_id == email_id)
    )
    existing_record = result.scalars().first()
    if existing_record:
        return {
            "rfq_id": existing_record.rfq_id,
            "status": existing_record.status,
            "email_id": email_id,
            "message": "RFQ already exists for this email message"
        }

    try:
        from services.gmail_inbox import get_email as _get
        em = await asyncio.to_thread(_get, email_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    import uuid
    rfq_id = str(uuid.uuid4())
    saved_excel_path = None

    # Extract text from PDF/Excel attachments and append to body
    attachment_texts = []
    try:
        from services.gmail_inbox import get_email_attachments_with_bytes
        from services.excel_extractor import extract_excel_bytes
        import io, pdfplumber

        SUPPORTED = {".pdf", ".xlsx", ".xlsm", ".xls", ".csv", ".docx", ".doc", ".txt"}
        attachments = await asyncio.to_thread(get_email_attachments_with_bytes, email_id)
        for att in attachments:
            fname = att["filename"]
            suffix = fname.rsplit(".", 1)[-1].lower() if "." in fname else ""
            if f".{suffix}" not in SUPPORTED:
                continue
            try:
                if suffix == "pdf":
                    with pdfplumber.open(io.BytesIO(att["bytes"])) as pdf:
                        pages = []
                        for page in pdf.pages:
                            t = page.extract_text()
                            if t:
                                pages.append(t)
                            for row in page.extract_tables():
                                pages.append(" | ".join(str(c or "") for c in row))
                        attachment_texts.append(f"[Attachment: {fname}]\n" + "\n".join(pages))
                elif suffix in ("xlsx", "xlsm", "xls", "csv"):
                    text = extract_excel_bytes(att["bytes"], fname)
                    attachment_texts.append(f"[Attachment: {fname}]\n{text}")
                    try:
                        import os
                        storage_dir = os.path.join(os.getcwd(), "storage", "attachments")
                        os.makedirs(storage_dir, exist_ok=True)
                        saved_excel_path = os.path.join(storage_dir, f"{rfq_id}_{fname}")
                        with open(saved_excel_path, "wb") as f:
                            f.write(att["bytes"])
                        logger.info(f"[EMAIL] Saved incoming pricing template to {saved_excel_path}")
                    except Exception as storage_ex:
                        logger.warning(f"[EMAIL] Could not save Excel attachment locally: {storage_ex}")
                elif suffix in ("docx", "doc"):
                    import docx, io as _io
                    doc = docx.Document(_io.BytesIO(att["bytes"]))
                    parts = [p.text for p in doc.paragraphs if p.text.strip()]
                    for table in doc.tables:
                        for row in table.rows:
                            parts.append(" | ".join(c.text for c in row.cells))
                    attachment_texts.append(f"[Attachment: {fname}]\n" + "\n".join(parts))
                elif suffix == "txt":
                    attachment_texts.append(f"[Attachment: {fname}]\n" + att["bytes"].decode("utf-8", errors="replace"))
            except Exception as ex:
                logger.warning(f"[EMAIL] Could not extract {fname}: {ex}")
    except Exception as ex:
        logger.warning(f"[EMAIL] Attachment extraction failed: {ex}")

    combined_body = em.get("body", "") or ""
    if attachment_texts:
        combined_body = combined_body + "\n\n" + "\n\n".join(attachment_texts)

    if not combined_body.strip():
        raise HTTPException(status_code=422, detail="Email has no text body or readable attachments")

    from workflow.graph import graph
    from db.crud import save_rfq

    subject = em.get("subject", "")
    sender = em.get("from_email", "")

    async def run_pipeline():
        try:
            from db.database import SessionLocal
            config = {"configurable": {"thread_id": rfq_id}}
            initial_state = {
                "rfq_id": rfq_id,
                "raw_content": combined_body,
                "subject": subject,
                "sender": sender,
                "source": "email",
                "email_id": email_id,
                "thread_id": em.get("thread_id"),
                "attachment_path": saved_excel_path,
            }
            result = await graph.ainvoke(initial_state, config=config)
            async with SessionLocal() as _db:
                await save_rfq(_db, result, result.get("status", "quoted"))
        except Exception as exc:
            logger.error(f"[EMAIL PIPELINE] {rfq_id}: {exc}", exc_info=True)

    background_tasks.add_task(run_pipeline)
    return {"rfq_id": rfq_id, "status": "processing", "email_id": email_id}
