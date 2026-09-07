import io
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border

logger = logging.getLogger(__name__)

def _copy_style(src, dest):
    if src.font:
        dest.font = Font(
            name=src.font.name,
            size=src.font.size,
            bold=src.font.bold,
            italic=src.font.italic,
            color=src.font.color,
        )
    if src.fill:
        dest.fill = PatternFill(
            fill_type=src.fill.fill_type,
            start_color=src.fill.start_color,
            end_color=src.fill.end_color,
        )
    if src.alignment:
        dest.alignment = Alignment(
            horizontal=src.alignment.horizontal,
            vertical=src.alignment.vertical,
            text_rotation=src.alignment.text_rotation,
            wrap_text=src.alignment.wrap_text,
            shrink_to_fit=src.alignment.shrink_to_fit,
            indent=src.alignment.indent,
        )
    if src.border:
        dest.border = Border(
            left=src.border.left,
            right=src.border.right,
            top=src.border.top,
            bottom=src.border.bottom,
        )

def fill_excel_template(file_path: str, line_pricing: list[dict]) -> bytes:
    """
    Reads the Excel file at file_path, fills calculated unit prices and subtotals
    matching by line item index, and returns the modified workbook bytes.
    """
    logger.info(f"[EXCEL FILLER] Filling pricing template at: {file_path}")
    try:
        wb = openpyxl.load_workbook(file_path)
        # Select active or search for common sheet names
        ws = wb.active
        for sheet_name in ["Pricing Template", "BOQ", "Bid Sheet"]:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                break

        # Map pricing by index
        pricing_map = {p["line_item_index"]: p for p in line_pricing}

        # Analyze headers in row 1
        max_col = ws.max_column
        headers = [str(ws.cell(row=1, column=c).value or "").lower() for c in range(1, max_col + 1)]

        unit_price_col = None
        subtotal_col = None

        # Look for explicit price/amount columns
        for col_idx, h in enumerate(headers, 1):
            if "unit price" in h or "unit rate" in h or ("rate" in h and "target" not in h) or ("price" in h and "total" not in h):
                unit_price_col = col_idx
            elif "subtotal" in h or "amount" in h or "line total" in h or "total price" in h:
                subtotal_col = col_idx

        # Fallback to remarks column if no price column is found
        if not unit_price_col:
            for col_idx, h in enumerate(headers, 1):
                if "remarks" in h or "comment" in h:
                    unit_price_col = col_idx
                    break

        # If still not found, append new columns
        if not unit_price_col:
            unit_price_col = max_col + 1
            ws.cell(row=1, column=unit_price_col, value="Quoted Unit Price (INR)")
            if max_col > 0:
                _copy_style(ws.cell(row=1, column=max_col), ws.cell(row=1, column=unit_price_col))
            max_col += 1

        if not subtotal_col:
            subtotal_col = unit_price_col + 1
            ws.cell(row=1, column=subtotal_col, value="Quoted Subtotal (INR)")
            _copy_style(ws.cell(row=1, column=unit_price_col), ws.cell(row=1, column=subtotal_col))

        # Fill values starting from row 2
        for r in range(2, ws.max_row + 1):
            # Row index matches pricing index (r - 2)
            idx = r - 2
            pricing = pricing_map.get(idx)
            if not pricing:
                continue

            u_cell = ws.cell(row=r, column=unit_price_col, value=pricing.get("unit_price", 0.0))
            s_cell = ws.cell(row=r, column=subtotal_col, value=pricing.get("subtotal", 0.0))

            # Style filled values
            u_cell.font = Font(name="Arial", size=10)
            u_cell.alignment = Alignment(horizontal="right")
            s_cell.font = Font(name="Arial", size=10, bold=True)
            s_cell.alignment = Alignment(horizontal="right")

        # Save workbook to memory buffer
        buf = io.BytesIO()
        wb.save(buf)
        logger.info("[EXCEL FILLER] Template successfully filled")
        return buf.getvalue()

    except Exception as e:
        logger.error(f"[EXCEL FILLER] Failed to fill pricing template: {e}", exc_info=True)
        # Return empty bytes so dispatch fallback can continue
        return b""
