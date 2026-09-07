import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(file_path)
    elif suffix in (".docx", ".doc"):
        return _extract_docx(file_path)
    elif suffix == ".txt":
        return path.read_text(encoding="utf-8")
    elif suffix in (".xlsx", ".xlsm", ".xls", ".csv"):
        from services.excel_extractor import extract_excel
        return extract_excel(file_path)
    elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
        from services.vision_extractor import extract_from_path
        return extract_from_path(file_path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


_VISION_THRESHOLD = 100  # chars — below this, treat PDF as scanned/image-based


def _extract_pdf(file_path: str) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
            for table in page.extract_tables():
                for row in table:
                    text_parts.append(" | ".join(str(cell or "") for cell in row))
    text = "\n".join(text_parts).strip()

    if len(text) < _VISION_THRESHOLD:
        logger.info(f"[PDF] Sparse text ({len(text)} chars) — falling back to Gemini Vision OCR")
        try:
            from services.vision_extractor import extract_from_path
            return extract_from_path(file_path)
        except Exception as e:
            logger.warning(f"[PDF] Vision fallback failed: {e} — returning partial text")

    return text


def _extract_docx(file_path: str) -> str:
    import docx
    doc = docx.Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
